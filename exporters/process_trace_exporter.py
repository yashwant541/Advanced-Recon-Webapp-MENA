"""Process-trace export: the "middle process" behind a completed run (spec:
"show the middle process - how things work").

Purpose
-------
Give a reviewer a workbook that answers "how did the engine get from these
uploaded files to that result" -- the standardized trial balance and
financial statement as actually used, the mapping table that was loaded,
every mapping-validation finding, mapping coverage, and a full per-line
calculation breakdown (formula, included/deducted/excluded accounts with
reasons, and any sign/unit warnings) for every calculator that ran. This is
strictly a read of data already captured on a completed
:class:`ReconciliationRun` -- it never reruns any calculation.

Public contents
----------------
``build_process_trace_workbook(run)`` -- returns an ``openpyxl.Workbook``.
``export_process_trace_bytes(run)`` -- returns ``.xlsx`` bytes.

Dependencies: ``exporters.excel_exporter``, ``exporters.standardized_exporter``.
"""

from __future__ import annotations

from openpyxl import Workbook

from iraq_recon.exporters.excel_exporter import add_table_sheet, export_workbook_bytes, new_workbook
from iraq_recon.exporters.standardized_exporter import financial_statement_rows, trial_balance_rows
from iraq_recon.models.reconciliation import ReconciliationRun


def _add_standardized_trial_balance(workbook: Workbook, run: ReconciliationRun) -> None:
    add_table_sheet(
        workbook, "Standardized Trial Balance",
        ["Account Number", "Account Description", "Row Type", "Posting Status", "Statement Type",
         "Natural Side", "Original Balance", "Source Unit", "Converted Balance", "Source File",
         "Source Sheet", "Source Row"],
        trial_balance_rows(list(run.trial_balance)),
        currency_columns=("Original Balance", "Converted Balance"),
        column_widths={"Account Description": 32},
    )


def _add_standardized_financial_statement(workbook: Workbook, run: ReconciliationRun) -> None:
    add_table_sheet(
        workbook, "Standardized Submission",
        ["Line Code", "Line Description", "Reported Amount", "Unit", "Currency Classification",
         "Residency Classification", "Schedule", "Source File", "Source Sheet", "Source Location"],
        financial_statement_rows(list(run.financial_statement)),
        currency_columns=("Reported Amount",),
        column_widths={"Line Description": 32},
    )


def _add_mapping_table(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Local Account", "Description", "Statement Type", "Natural Side", "Economic Role",
        "Ayra Category", "Iraq Reporting Bucket", "Financial Statement Line", "Schedule",
        "Presentation Sign", "Inclusion Status", "Mapping Method", "Mapping Confidence", "Rationale",
    ]
    rows = [
        [
            m.local_account, m.local_description, str(m.statement_type), str(m.natural_side),
            str(m.economic_role), m.ayra_category, m.iraq_reporting_bucket, m.financial_statement_line,
            m.schedule_code or "", m.presentation_sign, str(m.inclusion_status), str(m.mapping_method),
            m.mapping_confidence, m.mapping_rationale,
        ]
        for m in run.mapping_records
    ]
    add_table_sheet(
        workbook, "Mapping Table", headers, rows,
        column_widths={"Description": 30, "Rationale": 35},
    )


def _add_mapping_validation(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = ["Issue Code", "Severity", "Local Account", "Description"]
    rows = [
        [issue.issue_code, str(issue.severity), issue.local_account or "", issue.description]
        for issue in run.mapping_validation_issues
    ]
    add_table_sheet(
        workbook, "Mapping Validation", headers, rows,
        status_columns=("Severity",),
        column_widths={"Description": 45},
    )


def _add_mapping_coverage(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = ["Metric", "Value"]
    rows: list[list[object]] = []
    for key, value in run.mapping_coverage_summary.items():
        if isinstance(value, (list, tuple)):
            value = ", ".join(str(v) for v in value)
        elif isinstance(value, dict):
            value = str(value)
        rows.append([key, value])
    add_table_sheet(workbook, "Mapping Coverage", headers, rows, column_widths={"Metric": 28, "Value": 40})


def _add_calculation_detail(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Line Code", "Formula", "Calculated Amount", "Included Accounts", "Deducted Accounts",
        "Excluded Accounts (Reason)", "Warnings",
    ]
    rows = [
        [
            calc.line_code,
            calc.formula,
            calc.calculated_amount,
            ", ".join(calc.included_accounts),
            ", ".join(calc.deducted_accounts),
            "; ".join(f"{e.get('account', '')}: {e.get('reason', '')}" for e in calc.excluded_accounts),
            "; ".join(calc.warnings),
        ]
        for calc in run.calculation_results
    ]
    add_table_sheet(
        workbook, "Calculation Detail", headers, rows,
        currency_columns=("Calculated Amount",),
        column_widths={"Formula": 30, "Included Accounts": 30, "Excluded Accounts (Reason)": 45, "Warnings": 35},
    )


def _add_calculation_contributions(workbook: Workbook, run: ReconciliationRun) -> None:
    """Every account-level contribution behind every calculator, before
    aggregation into the final line result -- the finest-grained "how did
    we get this number" view."""
    headers = [
        "Line Code", "Account", "Description", "Formula Component", "Converted Amount",
        "Presentation Sign", "Presented Amount",
    ]
    rows = [
        [
            calc.line_code, c.account, c.description, c.formula_component,
            c.converted_amount, c.presentation_sign, c.presented_amount,
        ]
        for calc in run.calculation_results
        for c in calc.contributions
    ]
    add_table_sheet(
        workbook, "Calculation Contributions", headers, rows,
        currency_columns=("Converted Amount", "Presented Amount"),
        column_widths={"Description": 30},
    )


def build_process_trace_workbook(run: ReconciliationRun) -> Workbook:
    """Build the "middle process" workbook for a completed run.

    Args:
        run: A completed :class:`ReconciliationRun`.

    Returns:
        An ``openpyxl.Workbook`` with: Standardized Trial Balance,
        Standardized Submission, Mapping Table, Mapping Validation, Mapping
        Coverage, Calculation Detail, and Calculation Contributions.
    """
    workbook = new_workbook()
    _add_standardized_trial_balance(workbook, run)
    _add_standardized_financial_statement(workbook, run)
    _add_mapping_table(workbook, run)
    _add_mapping_validation(workbook, run)
    _add_mapping_coverage(workbook, run)
    _add_calculation_detail(workbook, run)
    _add_calculation_contributions(workbook, run)
    return workbook


def export_process_trace_bytes(run: ReconciliationRun) -> bytes:
    """Build the process-trace workbook and serialize it to ``.xlsx`` bytes."""
    return export_workbook_bytes(build_process_trace_workbook(run))
