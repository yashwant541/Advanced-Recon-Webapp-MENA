"""Audit-ready 14-sheet Excel export of a completed reconciliation run (spec section 20).

Purpose
-------
Build the full audit workbook from an already-completed
:class:`ReconciliationRun` -- Executive Summary, Statement Reconciliation,
Account Trace, Clubbing Breakdown, Schedule Tie-Out, Mapping Validation,
Unmapped Accounts, Exceptions, Control Results, Snapshot Comparison,
Excluded Rows, OFF BS Reconciliation, Configuration, and Run Metadata. This
module never reruns the reconciliation; it only reads what is already on
``run``.

Note on sheets 6, 7, and 11 ("Mapping Validation", "Unmapped Accounts",
"Excluded Rows"): :class:`ReconciliationRun` does not separately retain the
raw mapping-validation issue list or per-calculator excluded-account
reasons (those are intermediate engine state, not part of the run's public
result shape). These three sheets are therefore populated from ``run.exceptions``,
filtered by the diagnostic root-cause codes each concern maps to -- every
mapping/exclusion finding that reached the run's diagnostics is still
visible, just routed through the single exceptions list rather than a
second, separately-stored copy.

Public contents
----------------
``build_audit_workbook(run)`` -- returns an ``openpyxl.Workbook``.
``export_audit_workbook_bytes(run)`` -- returns ``.xlsx`` bytes.

Dependencies: ``exporters.excel_exporter``, ``exporters.formatting``.
"""

from __future__ import annotations

from decimal import Decimal
from typing import Any

from openpyxl import Workbook

from iraq_recon.constants import DiagnosticCode
from iraq_recon.exporters.excel_exporter import add_table_sheet, export_workbook_bytes, new_workbook
from iraq_recon.models.reconciliation import ReconciliationRun

_MAPPING_VALIDATION_CODES = frozenset(
    {
        DiagnosticCode.DUPLICATE_MAPPING,
        DiagnosticCode.CONFLICTING_MAPPING,
        DiagnosticCode.MISSING_PRESENTATION_SIGN,
        DiagnosticCode.MISSING_CONTRA_ACCOUNT,
    }
)
_EXCLUSION_CODES = frozenset(
    {
        DiagnosticCode.NON_POSTING_ROW_INCLUDED,
        DiagnosticCode.PARENT_CHILD_DOUBLE_COUNT,
        DiagnosticCode.OFF_BS_MISCLASSIFICATION,
        DiagnosticCode.DUPLICATE_SOURCE_ROW,
    }
)


def _exception_counts_by_line(run: ReconciliationRun) -> dict[str, int]:
    counts: dict[str, int] = {}
    for exc in run.exceptions:
        if exc.affected_line:
            counts[exc.affected_line] = counts.get(exc.affected_line, 0) + 1
    return counts


def _add_executive_summary(workbook: Workbook, run: ReconciliationRun) -> None:
    exception_counts = _exception_counts_by_line(run)
    headers = [
        "Reporting Line", "Reported Amount", "Calculated Amount", "Variance", "Status",
        "Supporting Account Count", "Supporting Schedule", "Exception Count", "Explanation",
    ]
    rows = [
        [
            line.line_code,
            line.reported_amount,
            line.calculated_amount,
            line.variance,
            str(line.status),
            len(line.supporting_accounts),
            line.supporting_schedule or "",
            exception_counts.get(line.line_code, 0),
            line.accounting_explanation,
        ]
        for line in run.line_results
    ]
    add_table_sheet(
        workbook, "Executive Summary", headers, rows,
        currency_columns=("Reported Amount", "Calculated Amount", "Variance"),
        status_columns=("Status",),
        column_widths={"Reporting Line": 28, "Explanation": 50},
    )


def _add_statement_reconciliation(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Line Code", "Line Description", "Reported Amount", "Calculated Amount", "Variance",
        "Absolute Variance", "Variance %", "Tolerance", "Status", "Formula", "Review Required",
    ]
    rows = [
        [
            line.line_code, line.line_description, line.reported_amount, line.calculated_amount,
            line.variance, line.absolute_variance,
            float(line.variance_percentage) if line.variance_percentage is not None else None,
            line.tolerance, str(line.status), line.formula, "YES" if line.review_required else "NO",
        ]
        for line in run.line_results
    ]
    add_table_sheet(
        workbook, "Statement Reconciliation", headers, rows,
        currency_columns=("Reported Amount", "Calculated Amount", "Variance", "Absolute Variance", "Tolerance"),
        percentage_columns=("Variance %",),
        status_columns=("Status",),
        column_widths={"Line Description": 32, "Formula": 40},
    )


def _add_account_trace(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Source File", "Source Sheet", "Source Row", "Account", "Description", "Original Balance",
        "Source Unit", "Conversion Factor", "Converted Balance", "Statement Type", "Natural Side",
        "Economic Role", "Ayra Category", "Iraq Reporting Bucket", "Financial Statement Line",
        "Schedule", "Presentation Sign", "Presented Amount", "Mapping Method", "Mapping Confidence",
        "Rationale", "Review Status",
    ]
    rows = [
        [
            e.source_file, e.source_sheet, e.source_row, e.account, e.description, e.original_balance,
            e.source_unit, e.conversion_factor, e.converted_balance, e.statement_type, e.natural_side,
            e.economic_role, e.ayra_category, e.iraq_reporting_bucket, e.financial_statement_line,
            e.schedule_code or "", e.presentation_sign, e.presented_amount, e.mapping_method,
            e.mapping_confidence, e.mapping_rationale, e.review_status,
        ]
        for e in run.account_trace
    ]
    add_table_sheet(
        workbook, "Account Trace", headers, rows,
        currency_columns=("Original Balance", "Converted Balance", "Presented Amount"),
        column_widths={"Description": 32, "Rationale": 40},
    )


def _add_clubbing_breakdown(workbook: Workbook, run: ReconciliationRun) -> None:
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for e in run.account_trace:
        key = (e.iraq_reporting_bucket, e.ayra_category)
        group = grouped.setdefault(key, {"total": Decimal("0"), "accounts": []})
        group["total"] += e.presented_amount
        group["accounts"].append(e.account)

    headers = ["Iraq Reporting Bucket", "Ayra Category", "Total", "Account Count", "Included Accounts"]
    rows = [
        [bucket, category, group["total"], len(group["accounts"]), ", ".join(group["accounts"])]
        for (bucket, category), group in grouped.items()
    ]
    add_table_sheet(
        workbook, "Clubbing Breakdown", headers, rows,
        currency_columns=("Total",),
        column_widths={"Included Accounts": 50},
    )


def _add_schedule_tie_out(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Schedule Code", "Description", "Schedule Total", "TB Total", "Statement Total",
        "Variance to TB", "Variance to Statement", "Status",
    ]
    rows = [
        [
            s.schedule_code, s.schedule_description, s.schedule_total, s.tb_total, s.statement_total,
            s.variance_to_tb, s.variance_to_statement, str(s.status),
        ]
        for s in run.schedule_results
    ]
    add_table_sheet(
        workbook, "Schedule Tie-Out", headers, rows,
        currency_columns=("Schedule Total", "TB Total", "Statement Total", "Variance to TB", "Variance to Statement"),
        status_columns=("Status",),
        column_widths={"Description": 32},
    )


def _add_exception_derived_sheet(
    workbook: Workbook, run: ReconciliationRun, title: str, codes: frozenset[DiagnosticCode]
) -> None:
    headers = ["Root Cause", "Affected Account", "Affected Line", "Severity", "Likely Cause", "Suggested Review Action"]
    rows = [
        [str(e.root_cause_code), e.affected_account or "", e.affected_line or "", str(e.severity),
         e.likely_cause, e.suggested_review_action]
        for e in run.exceptions
        if e.root_cause_code in codes
    ]
    add_table_sheet(
        workbook, title, headers, rows,
        status_columns=("Severity",),
        column_widths={"Likely Cause": 45, "Suggested Review Action": 45},
    )


def _add_unmapped_accounts(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = ["Account", "Severity", "Balance", "Likely Cause", "Suggested Review Action"]
    rows = [
        [
            e.affected_account or "", str(e.severity), e.supporting_values.get("balance", ""),
            e.likely_cause, e.suggested_review_action,
        ]
        for e in run.exceptions
        if e.root_cause_code == DiagnosticCode.UNMAPPED_ACCOUNT
    ]
    add_table_sheet(
        workbook, "Unmapped Accounts", headers, rows,
        status_columns=("Severity",),
        column_widths={"Likely Cause": 45, "Suggested Review Action": 45},
    )


def _add_exceptions(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = [
        "Root Cause", "Affected Account", "Affected Line", "Severity", "Can Continue",
        "Likely Cause", "Suggested Review Action",
    ]
    rows = [
        [
            str(e.root_cause_code), e.affected_account or "", e.affected_line or "", str(e.severity),
            "YES" if e.can_continue else "NO", e.likely_cause, e.suggested_review_action,
        ]
        for e in run.exceptions
    ]
    add_table_sheet(
        workbook, "Exceptions", headers, rows,
        status_columns=("Severity",),
        column_widths={"Likely Cause": 45, "Suggested Review Action": 45},
    )


def _add_control_results(workbook: Workbook, run: ReconciliationRun) -> None:
    headers = ["Control Code", "Description", "Status", "Severity", "Expected", "Actual", "Variance"]
    rows = [
        [
            c.control_code, c.control_description, str(c.status), str(c.severity),
            c.expected_amount, c.actual_amount, c.variance,
        ]
        for c in run.control_results
    ]
    add_table_sheet(
        workbook, "Control Results", headers, rows,
        currency_columns=("Expected", "Actual", "Variance"),
        status_columns=("Status",),
        column_widths={"Description": 45},
    )


def _add_snapshot_comparison(workbook: Workbook, run: ReconciliationRun) -> None:
    candidates = run.candidate_snapshot_comparison.get("candidates", [])
    headers = ["Snapshot", "Matched Anchors", "Exact Matches", "Precision Matches", "Total Abs Variance", "Value Coverage"]
    rows = [
        [
            c.get("snapshot_id", ""), c.get("matched_anchor_count", ""), c.get("exact_match_count", ""),
            c.get("precision_match_count", ""), c.get("total_absolute_variance", ""),
            c.get("mapping_value_coverage", ""),
        ]
        for c in candidates
    ]
    add_table_sheet(workbook, "Snapshot Comparison", headers, rows)


def _add_off_balance_sheet(workbook: Workbook, run: ReconciliationRun) -> None:
    obs_lines = [line for line in run.line_results if "OFF_BALANCE_SHEET" in line.line_code]
    obs_controls = [c for c in run.control_results if c.control_code.startswith("OBS_")]

    headers = ["Line Code", "Reported Amount", "Calculated Amount", "Variance", "Status"]
    rows = [
        [line.line_code, line.reported_amount, line.calculated_amount, line.variance, str(line.status)]
        for line in obs_lines
    ]
    rows.extend(
        [c.control_code, c.expected_amount, c.actual_amount, c.variance, str(c.status)] for c in obs_controls
    )
    add_table_sheet(
        workbook, "OFF BS Reconciliation", headers, rows,
        currency_columns=("Reported Amount", "Calculated Amount", "Variance"),
        status_columns=("Status",),
    )


def _flatten_config(prefix: str, value: Any, rows: list[list[str]]) -> None:
    if isinstance(value, dict):
        for key, sub_value in value.items():
            _flatten_config(f"{prefix}.{key}" if prefix else key, sub_value, rows)
    else:
        rows.append([prefix, str(value)])


def _add_configuration(workbook: Workbook, run: ReconciliationRun) -> None:
    rows: list[list[str]] = []
    _flatten_config("", run.configuration.to_dict(), rows)
    add_table_sheet(workbook, "Configuration", ["Setting", "Value"], rows, column_widths={"Setting": 40, "Value": 40})


def _add_run_metadata(workbook: Workbook, run: ReconciliationRun) -> None:
    rows = [
        ["Run ID", run.run_id],
        ["Status", run.status],
        ["Selected TB Snapshot", run.selected_tb_snapshot],
        ["Line Count", str(len(run.line_results))],
        ["Schedule Count", str(len(run.schedule_results))],
        ["Control Count", str(len(run.control_results))],
        ["Exception Count", str(len(run.exceptions))],
        ["Warnings", "; ".join(run.warnings)],
    ]
    add_table_sheet(workbook, "Run Metadata", ["Field", "Value"], rows, column_widths={"Field": 24, "Value": 60})


def build_audit_workbook(run: ReconciliationRun) -> Workbook:
    """Build the full 14-sheet audit workbook from a completed run.

    Args:
        run: A completed :class:`ReconciliationRun`.

    Returns:
        An ``openpyxl.Workbook`` with every spec section 20 sheet.
    """
    workbook = new_workbook()
    _add_executive_summary(workbook, run)
    _add_statement_reconciliation(workbook, run)
    _add_account_trace(workbook, run)
    _add_clubbing_breakdown(workbook, run)
    _add_schedule_tie_out(workbook, run)
    _add_exception_derived_sheet(workbook, run, "Mapping Validation", _MAPPING_VALIDATION_CODES)
    _add_unmapped_accounts(workbook, run)
    _add_exceptions(workbook, run)
    _add_control_results(workbook, run)
    _add_snapshot_comparison(workbook, run)
    _add_exception_derived_sheet(workbook, run, "Excluded Rows", _EXCLUSION_CODES)
    _add_off_balance_sheet(workbook, run)
    _add_configuration(workbook, run)
    _add_run_metadata(workbook, run)
    return workbook


def export_audit_workbook_bytes(run: ReconciliationRun) -> bytes:
    """Build the audit workbook and serialize it to ``.xlsx`` bytes."""
    return export_workbook_bytes(build_audit_workbook(run))
