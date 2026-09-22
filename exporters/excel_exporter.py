"""Generic Excel workbook-building mechanics.

Purpose
-------
Provide the reusable, sheet-agnostic mechanics for building a multi-sheet
audit workbook: creating a workbook, adding one formatted table sheet at a
time, and serializing the result to bytes. ``exporters.audit_exporter``
uses these helpers to assemble the specific 14 spec-mandated sheets from a
:class:`ReconciliationRun`; this module knows nothing about reconciliation
domain concepts.

Public contents
----------------
``new_workbook()`` -- an empty workbook with no default sheet.
``add_table_sheet(workbook, title, headers, rows, ...)`` -- add one styled sheet.
``export_workbook_bytes(workbook)`` -- serialize to ``.xlsx`` bytes.

Dependencies: ``openpyxl``, ``exporters.formatting``.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from iraq_recon.exceptions import ExportError
from iraq_recon.exporters.formatting import (
    CURRENCY_FORMAT,
    PERCENTAGE_FORMAT,
    apply_column_widths,
    apply_number_format,
    fill_status_column,
    freeze_and_filter,
    style_header_row,
)

#: Excel worksheet titles cannot exceed 31 characters.
_MAX_SHEET_TITLE_LENGTH = 31


def new_workbook() -> Workbook:
    """Create an empty workbook with its default sheet removed."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    return workbook


def add_table_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    currency_columns: tuple[str, ...] = (),
    percentage_columns: tuple[str, ...] = (),
    status_columns: tuple[str, ...] = (),
    column_widths: dict[str, int] | None = None,
) -> Worksheet:
    """Add one styled table sheet to a workbook.

    Args:
        workbook: The workbook to add to.
        title: Sheet title (truncated to Excel's 31-character limit).
        headers: Column headers, written to row 1.
        rows: Data rows; each must have the same length as ``headers``.
        currency_columns: Header names to format as currency.
        percentage_columns: Header names to format as percentages.
        status_columns: Header names to color-fill by status value.
        column_widths: ``{header_name: width}`` for explicit column sizing.

    Returns:
        The created :class:`openpyxl.worksheet.worksheet.Worksheet`.

    Raises:
        ExportError: if any row's length does not match ``headers``.
    """
    for row in rows:
        if len(row) != len(headers):
            raise ExportError(
                f"Row length {len(row)} does not match header length {len(headers)} on sheet '{title}'.",
                details={"sheet": title, "headers": headers, "row": row},
            )

    worksheet = workbook.create_sheet(title=title[:_MAX_SHEET_TITLE_LENGTH])
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    style_header_row(worksheet)

    header_to_letter = {header: get_column_letter(index + 1) for index, header in enumerate(headers)}

    for header in currency_columns:
        if header in header_to_letter:
            apply_number_format(worksheet, header_to_letter[header], CURRENCY_FORMAT)
    for header in percentage_columns:
        if header in header_to_letter:
            apply_number_format(worksheet, header_to_letter[header], PERCENTAGE_FORMAT)
    for header in status_columns:
        if header in header_to_letter:
            fill_status_column(worksheet, header_to_letter[header])

    if column_widths:
        apply_column_widths(
            worksheet, {header_to_letter[h]: w for h, w in column_widths.items() if h in header_to_letter}
        )

    freeze_and_filter(worksheet)
    return worksheet


def export_workbook_bytes(workbook: Workbook) -> bytes:
    """Serialize a workbook to ``.xlsx`` bytes."""
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
